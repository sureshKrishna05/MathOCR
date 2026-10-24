"""Comprehensive tests for structured document and native OMML DOCX export in MathOCR.

Covers:
- Plain text paragraphs
- Native OMML fractions, superscripts/subscripts, integrals/summations, square roots
- Mixed paragraphs with inline equations and formatted runs
- Standalone display equations with center alignment
- Genuine Word tables with editable rows, columns, cells, and embedded math
- Mixed academic document (headings, text, equations, tables, lists)
- Equation conversion failure handling (graceful warning, zero LaTeX leakage)
- 100% offline local generation and LibreOffice headless verification
"""

import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from backend.app import app
from backend.document import (
    DisplayEquation,
    Document,
    Heading,
    InlineEquation,
    ListBlock,
    ListItem,
    Paragraph,
    Table,
    TableCell,
    TableRow,
    TextRun,
)
from backend.docx_export import build_docx_from_document, verify_docx_file
from backend.omml import OMML_NS, latex_to_omml, validate_equation

client = TestClient(app)


def _read_docx_xml(docx_path: Path) -> str:
    """Extract word/document.xml content as a string."""
    with zipfile.ZipFile(docx_path, "r") as zf:
        return zf.read("word/document.xml").decode("utf-8")


# 1. Plain text document
def test_plain_text_document_stored_as_editable_word_paragraphs(tmp_path: Path) -> None:
    doc = Document(
        title="Plain Text Document",
        elements=[
            Paragraph(runs=[TextRun(text="First paragraph of plain editable text.")]),
            Paragraph(runs=[TextRun(text="Second paragraph with regular body content.")]),
        ],
    )
    output = tmp_path / "plain_text.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0
    assert output.exists()

    xml = _read_docx_xml(output)
    assert "<w:p" in xml
    assert "First paragraph of plain editable text." in xml
    assert "Second paragraph with regular body content." in xml

    # Verify with python-docx parser
    loaded = docx.Document(str(output))
    para_texts = [p.text for p in loaded.paragraphs]
    assert any("First paragraph of plain editable text." in t for t in para_texts)
    assert any("Second paragraph with regular body content." in t for t in para_texts)


# 2. Fraction equation
def test_fraction_equation_creates_native_omml_fraction(tmp_path: Path) -> None:
    latex = r"\frac{a+b}{c}"
    doc = Document(
        title="Fraction Test",
        elements=[
            DisplayEquation(latex=latex),
        ],
    )
    output = tmp_path / "fraction.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    # Must use OMML namespace and native fraction element m:f
    assert OMML_NS in xml
    assert "m:oMath" in xml
    assert "m:f" in xml
    assert "m:num" in xml
    assert "m:den" in xml
    assert r"\frac" not in xml  # No raw LaTeX


# 3. Superscripts and subscripts
def test_superscripts_and_subscripts_creates_structured_math(tmp_path: Path) -> None:
    latex = r"x^2 + H_2O"
    doc = Document(
        title="SubSup Test",
        elements=[
            DisplayEquation(latex=latex),
        ],
    )
    output = tmp_path / "subsup.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert OMML_NS in xml
    assert "m:oMath" in xml
    # Check for superscript and subscript OMML tags
    assert "m:sSup" in xml or "m:sSub" in xml
    assert "^" not in xml or "m:sup" in xml


# 4. Complex mathematical expression
def test_complex_mathematical_expression_creates_native_equation_without_latex(
    tmp_path: Path,
) -> None:
    latex = r"\int_0^1 x^2 \, dx = \frac{1}{3}"
    doc = Document(
        title="Complex Math Test",
        elements=[
            DisplayEquation(latex=latex),
        ],
    )
    output = tmp_path / "complex.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert "m:oMath" in xml
    assert "m:nary" in xml  # Integral / n-ary operator
    assert "m:f" in xml  # Fraction
    assert r"\int" not in xml  # Never raw LaTeX
    assert r"\frac" not in xml


# 5. Mixed paragraph
def test_mixed_paragraph_inlines_equation_in_correct_order(tmp_path: Path) -> None:
    doc = Document(
        title="Mixed Paragraph Test",
        elements=[
            Paragraph(
                runs=[
                    TextRun(text="Energy is given by "),
                    InlineEquation(latex=r"E = mc^2"),
                    TextRun(text=", as proposed by Einstein."),
                ]
            )
        ],
    )
    output = tmp_path / "mixed_paragraph.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert "Energy is given by" in xml
    assert "as proposed by Einstein." in xml
    assert "m:oMath" in xml
    assert "m:sSup" in xml

    # Verify structure: runs and m:oMath are siblings inside one <w:p>
    loaded = docx.Document(str(output))
    found_mixed = False
    for p in loaded.paragraphs:
        if "Energy is given by" in p.text:
            found_mixed = True
            child_tags = [child.tag.split("}")[-1] for child in p._p]
            assert "r" in child_tags
            assert "oMath" in child_tags
    assert found_mixed


# 6. Display equation
def test_display_equation_placed_standalone_with_omml_para(tmp_path: Path) -> None:
    doc = Document(
        title="Display Equation Test",
        elements=[
            Paragraph(runs=[TextRun(text="Introductory text before equation:")]),
            DisplayEquation(latex=r"\sum_{n=1}^{\infty} \frac{1}{n^2} = \frac{\pi^2}{6}"),
            Paragraph(runs=[TextRun(text="Explanatory text after equation.")]),
        ],
    )
    output = tmp_path / "display_equation.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert "m:oMathPara" in xml
    assert "m:oMath" in xml
    assert "m:nary" in xml

    loaded = docx.Document(str(output))
    assert len(loaded.paragraphs) >= 3


# 7. Table reconstruction
def test_table_reconstruction_creates_actual_word_table_with_cells(tmp_path: Path) -> None:
    table = Table(
        rows=[
            TableRow(
                cells=[
                    TableCell(elements=[Paragraph(runs=[TextRun(text="Method", bold=True)])]),
                    TableCell(elements=[Paragraph(runs=[TextRun(text="Formula", bold=True)])]),
                    TableCell(elements=[Paragraph(runs=[TextRun(text="Score", bold=True)])]),
                ],
                is_header=True,
            ),
            TableRow(
                cells=[
                    TableCell(elements=[Paragraph(runs=[TextRun(text="Model A")])]),
                    TableCell(elements=[Paragraph(runs=[InlineEquation(latex=r"\alpha + \beta")])]),
                    TableCell(elements=[Paragraph(runs=[TextRun(text="98.5%")])]),
                ]
            ),
            TableRow(
                cells=[
                    TableCell(elements=[Paragraph(runs=[TextRun(text="Model B")])]),
                    TableCell(elements=[Paragraph(runs=[InlineEquation(latex=r"\sqrt{x^2+y^2}")])]),
                    TableCell(elements=[Paragraph(runs=[TextRun(text="99.1%")])]),
                ]
            ),
        ]
    )
    doc = Document(title="Table Test", elements=[table])
    output = tmp_path / "table_test.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert "<w:tbl" in xml
    assert "<w:tr" in xml
    assert "<w:tc" in xml
    assert "Model A" in xml
    assert "Model B" in xml
    assert "98.5%" in xml
    assert "m:oMath" in xml
    assert "m:rad" in xml

    # Verify table structure with python-docx
    loaded = docx.Document(str(output))
    assert len(loaded.tables) == 1
    t = loaded.tables[0]
    assert len(t.rows) == 3
    assert len(t.columns) == 3
    assert "Method" in t.cell(0, 0).text
    assert "Model A" in t.cell(1, 0).text


# 8. Mixed academic document
def test_mixed_academic_document_coexists_harmoniously(tmp_path: Path) -> None:
    doc = Document(
        title="Quantum Harmonic Oscillator Analysis",
        elements=[
            Heading(level=1, text="1. Introduction"),
            Paragraph(
                runs=[
                    TextRun(text="The time-independent Schrödinger equation with potential "),
                    InlineEquation(latex=r"V(x) = \frac{1}{2}m\omega^2 x^2"),
                    TextRun(text=" takes the form:"),
                ]
            ),
            DisplayEquation(latex=r"-\frac{\hbar^2}{2m} \frac{d^2\psi}{dx^2} + \frac{1}{2}m\omega^2 x^2 \psi = E\psi"),
            Heading(level=2, text="2. Energy Eigenvalues"),
            Paragraph(
                runs=[
                    TextRun(text="The exact eigenvalues are given by:"),
                ]
            ),
            DisplayEquation(latex=r"E_n = \left(n + \frac{1}{2}\right)\hbar\omega"),
            Table(
                rows=[
                    TableRow(
                        cells=[
                            TableCell(elements=[Paragraph(runs=[TextRun(text="Quantum State n")])]),
                            TableCell(elements=[Paragraph(runs=[TextRun(text="Energy En")])]),
                        ],
                        is_header=True,
                    ),
                    TableRow(
                        cells=[
                            TableCell(elements=[Paragraph(runs=[TextRun(text="0 (Ground state)")])]),
                            TableCell(elements=[Paragraph(runs=[InlineEquation(latex=r"\frac{1}{2}\hbar\omega")])]),
                        ]
                    ),
                ]
            ),
            ListBlock(
                items=[
                    ListItem(elements=[Paragraph(runs=[TextRun(text="Equally spaced levels")])]),
                    ListItem(elements=[Paragraph(runs=[TextRun(text="Non-zero ground state energy")])]),
                ],
                ordered=False,
            ),
        ],
    )
    output = tmp_path / "academic_full.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert "<w:tbl" in xml
    assert "m:oMath" in xml
    assert "Quantum Harmonic Oscillator" in xml
    assert "1. Introduction" in xml

    # Verify with LibreOffice headless conversion to PDF if LibreOffice is present
    if shutil.which("libreoffice"):
        cmd = ["libreoffice", "--headless", "--convert-to", "pdf", str(output), "--outdir", str(tmp_path)]
        res = subprocess.run(cmd, capture_output=True, text=True)
        assert res.returncode == 0
        pdf_file = tmp_path / "academic_full.pdf"
        assert pdf_file.exists() and pdf_file.stat().st_size > 0


# 9. Conversion failure
def test_conversion_failure_produces_warning_and_no_visible_raw_latex(tmp_path: Path) -> None:
    doc = Document(
        title="Failure Test",
        elements=[
            Paragraph(
                runs=[
                    TextRun(text="Before bad math "),
                    InlineEquation(latex=r"\frac{broken}{"),
                    TextRun(text=" after bad math."),
                ]
            ),
            DisplayEquation(latex=r"\int_{0}^{1 broken"),
        ],
    )
    output = tmp_path / "failure_test.docx"
    warnings = build_docx_from_document(doc, output)

    # Must produce clear warnings
    assert len(warnings) >= 2
    assert any("could not be converted" in w for w in warnings)

    # Verify that raw LaTeX does NOT leak into document.xml text runs
    xml = _read_docx_xml(output)
    assert r"\frac{broken}{" not in xml
    assert r"\int_{0}^{1 broken" not in xml
    # Verify that a placeholder warning text exists
    assert "could not be converted" in xml


# 10. Offline export
def test_offline_export_works_without_network_access(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # Disable network calls by blocking socket
    import socket

    def guarded_connect(*args, **kwargs):
        raise OSError("Network access disabled during offline export test")

    monkeypatch.setattr(socket.socket, "connect", guarded_connect)

    doc = Document(
        title="Offline Test",
        elements=[
            Paragraph(runs=[TextRun(text="Offline document content with math: ")]),
            DisplayEquation(latex=r"\lim_{x \to 0} \frac{\sin x}{x} = 1"),
        ],
    )
    output = tmp_path / "offline_test.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0
    assert output.exists()

    xml = _read_docx_xml(output)
    assert "m:oMath" in xml
    assert "m:limLow" in xml


# 11. API endpoint export test
def test_api_export_docx_with_markdown_tables_and_math() -> None:
    markdown = """# API Test Document

Here is a paragraph with $E=mc^2$.

| Column 1 | Column 2 |
| :--- | :--- |
| Val A | $\\frac{x}{y}$ |

$$
\\int_0^1 x^2 \\, dx = \\frac{1}{3}
$$
"""
    response = client.post(
        "/api/export/docx",
        json={"title": "API Test", "markdown": markdown},
    )
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "<w:tbl" in xml
    assert "m:oMath" in xml
    assert "m:f" in xml
    assert "m:nary" in xml


def test_api_export_docx_with_structured_document() -> None:
    doc = Document(
        title="Structured API Test",
        elements=[
            Heading(level=1, text="Structured Heading"),
            Paragraph(runs=[TextRun(text="Structured paragraph with "), InlineEquation(latex=r"\sqrt{z}")]),
        ],
    )
    response = client.post(
        "/api/export/docx",
        json={"title": "Structured API", "document": doc.to_dict()},
    )
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "Structured Heading" in xml
    assert "m:oMath" in xml
    assert "m:rad" in xml


def test_empty_document_raises_422() -> None:
    response = client.post(
        "/api/export/docx",
        json={"title": "Empty Doc"},
    )
    assert response.status_code == 422


def test_multi_page_document_export(tmp_path: Path) -> None:
    from backend.document import PageBreak

    doc = Document(
        title="Multi-Page Document",
        elements=[
            Heading(level=1, text="Page One Content"),
            Paragraph(runs=[TextRun(text="Paragraph on page 1 with equation "), InlineEquation(latex=r"E_1 = \frac{1}{2}\hbar\omega")]),
            PageBreak(page=1),
            Heading(level=1, text="Page Two Content"),
            Paragraph(runs=[TextRun(text="Paragraph on page 2 with equation "), InlineEquation(latex=r"E_2 = \frac{3}{2}\hbar\omega")]),
        ],
    )
    output = tmp_path / "multipage.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    xml = _read_docx_xml(output)
    assert "w:br" in xml
    assert "Page One Content" in xml
    assert "Page Two Content" in xml
    assert "m:oMath" in xml
    assert "m:f" in xml


def test_end_to_end_ocr_to_docx_pipeline(tmp_path: Path) -> None:
    from PIL import Image, ImageDraw

    # Create synthetic academic page image with formula
    img = Image.new("RGB", (600, 200), color="white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 30), "Harmonic Oscillator Energy:", fill="black")
    draw.text((60, 80), "E = (n + 1/2) hbar omega", fill="black")

    buf = io.BytesIO()
    img.save(buf, format="PNG")

    # Step 1: Real OCR processing through /api/ocr endpoint
    files = [("files", ("harmonic_test.png", buf.getvalue(), "image/png"))]
    resp_ocr = client.post("/api/ocr", files=files, data={"mode": "formula", "engine": "auto"})
    assert resp_ocr.status_code == 200
    ocr_payload = resp_ocr.json()
    assert len(ocr_payload["results"]) >= 1
    recognized_latex = ocr_payload["results"][0]["latex"]
    assert len(recognized_latex) > 0

    # Step 2: Real DOCX export of recognized equation through /api/export/docx endpoint
    export_payload = {
        "title": "Harmonic Oscillator Analysis",
        "equations": [recognized_latex],
    }
    resp_docx = client.post("/api/export/docx", json=export_payload)
    assert resp_docx.status_code == 200

    # Step 3: Inspect generated DOCX XML
    docx_file = tmp_path / "end_to_end_result.docx"
    docx_file.write_bytes(resp_docx.content)
    xml = _read_docx_xml(docx_file)

    assert "m:oMath" in xml
    assert "<w:p" in xml
    assert "<w:r" in xml

    # Step 4: Verify with LibreOffice headless if available
    if shutil.which("libreoffice"):
        cmd = ["libreoffice", "--headless", "--convert-to", "pdf", str(docx_file), "--outdir", str(tmp_path)]
        res = subprocess.run(cmd, capture_output=True, text=True)
        assert res.returncode == 0
        pdf_file = tmp_path / "end_to_end_result.pdf"
        assert pdf_file.exists() and pdf_file.stat().st_size > 0


# 16. Regression: Markdown editor edits must override stale document dict
def test_markdown_editor_edits_override_stale_document_object(tmp_path: Path) -> None:
    stale_doc = Document(
        title="Stale Document",
        elements=[Paragraph(runs=[TextRun(text="OLD STALE UNEDITED OCR CONTENT")])],
    )
    edited_markdown = "# Edited Title\n\nNEW EDITED USER CONTENT in paragraph."

    response = client.post(
        "/api/export/docx",
        json={
            "title": "Edited Title",
            "markdown": edited_markdown,
            "document": stale_doc.to_dict(),
            "save_to_downloads": False,
        },
    )
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")

    assert "NEW EDITED USER CONTENT" in xml
    assert "OLD STALE UNEDITED OCR CONTENT" not in xml


# 17. Regression: Markdown page break preservation
def test_markdown_page_break_preservation(tmp_path: Path) -> None:
    markdown = "Page 1 Content\n\n---\n\nPage 2 Content"
    doc = Document.from_markdown(markdown)
    output = tmp_path / "page_break_test.docx"
    build_docx_from_document(doc, output)

    xml = _read_docx_xml(output)
    assert "w:br" in xml
    assert "Page 1 Content" in xml
    assert "Page 2 Content" in xml


# 18. ImageBlock embedding
def test_image_block_export(tmp_path: Path) -> None:
    from PIL import Image
    from backend.document import ImageBlock

    img_path = tmp_path / "test_figure.png"
    img = Image.new("RGB", (100, 100), color="blue")
    img.save(img_path)

    doc = Document(
        title="Image Test",
        elements=[
            Heading(level=1, text="Figure Section"),
            ImageBlock(src=str(img_path), caption="A test figure caption"),
        ],
    )
    output = tmp_path / "image_test.docx"
    warnings = build_docx_from_document(doc, output)
    assert len(warnings) == 0

    with zipfile.ZipFile(output) as zf:
        namelist = zf.namelist()
        assert any(n.startswith("word/media/") for n in namelist)
        xml = zf.read("word/document.xml").decode("utf-8")
        assert "Figure: A test figure caption" in xml


# 19. All four template presets verification
@pytest.mark.parametrize("preset", ["academic_textbook", "research_paper", "standard", "preserve_layout"])
def test_all_four_template_presets(tmp_path: Path, preset: str) -> None:
    doc = Document(
        title=f"Test {preset} Document",
        elements=[
            Heading(level=1, text=f"Heading in {preset}"),
            Paragraph(runs=[TextRun(text=f"Body paragraph for {preset} with equation "), InlineEquation(latex=r"\alpha + \beta = \gamma")]),
            DisplayEquation(latex=r"\int_0^\infty e^{-x^2} \, dx = \frac{\sqrt{\pi}}{2}"),
            Table(
                rows=[
                    TableRow(cells=[TableCell(elements=[Paragraph(runs=[TextRun(text="Col A", bold=True)])]), TableCell(elements=[Paragraph(runs=[TextRun(text="Col B", bold=True)])])], is_header=True),
                    TableRow(cells=[TableCell(elements=[Paragraph(runs=[TextRun(text="Val 1")])]), TableCell(elements=[Paragraph(runs=[InlineEquation(latex=r"\frac{1}{2}")])])]),
                ]
            ),
        ],
    )
    output = tmp_path / f"{preset}_output.docx"
    warnings = build_docx_from_document(doc, output, template=preset)
    assert len(warnings) == 0
    assert output.exists()

    with zipfile.ZipFile(output) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
        assert "m:oMath" in xml
        assert "<w:tbl" in xml

        # Verify footer contains page numbering
        footers = [f for f in zf.namelist() if f.startswith("word/footer")]
        assert len(footers) > 0
        footer_xml = zf.read(footers[0]).decode("utf-8")
        assert 'w:instr="PAGE"' in footer_xml

    # Verify LibreOffice can convert it to PDF
    if shutil.which("libreoffice"):
        cmd = ["libreoffice", "--headless", "--convert-to", "pdf", str(output), "--outdir", str(tmp_path)]
        res = subprocess.run(cmd, capture_output=True, text=True)
        assert res.returncode == 0
        pdf_file = tmp_path / f"{preset}_output.pdf"
        assert pdf_file.exists() and pdf_file.stat().st_size > 0


# 20. Real frontend API binary download flow
def test_api_binary_docx_download_headers() -> None:
    payload = {
        "title": "Quantum Physics Notes",
        "markdown": "# Quantum Physics\n\nEnergy relation:\n\n$$\nE = h\\nu\n$$\n",
        "template": "academic_textbook",
        "save_to_downloads": False,
    }
    response = client.post("/api/export/docx", json=payload)
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers["content-type"]
    assert "attachment;" in response.headers["content-disposition"]
    assert ".docx" in response.headers["content-disposition"]

    # Verify response body is a valid DOCX ZIP package
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        namelist = zf.namelist()
        assert "[Content_Types].xml" in namelist
        assert "word/document.xml" in namelist
        xml = zf.read("word/document.xml").decode("utf-8")
        assert "m:oMath" in xml


# 21. Error handling does not return binary
def test_api_error_handling_does_not_return_binary() -> None:
    # Empty request
    response = client.post("/api/export/docx", json={"title": "Empty", "save_to_downloads": False})
    assert response.status_code == 422
    assert "application/json" in response.headers["content-type"]
    err = response.json()
    assert "detail" in err

